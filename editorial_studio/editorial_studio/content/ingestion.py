from __future__ import annotations
import re
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Optional
import json

from editorial_studio.core.models import (
    ContentBlock,
    ContentType,
    Manuscript,
    SemanticRole,
)


@dataclass
class IngestionResult:
    manuscript: Manuscript
    warnings: list[str]
    errors: list[str]
    stats: dict[str, Any]


class ManuscriptIngester:
    def __init__(self):
        self.block_counter = 0
        self.warnings: list[str] = []
        self.errors: list[str] = []

    def ingest_file(self, file_path: str | Path, title: str = "", author: str = "") -> IngestionResult:
        path = Path(file_path)
        if not path.exists():
            self.errors.append(f"File not found: {file_path}")
            return IngestionResult(
                manuscript=Manuscript(id="", title=title or path.stem, author=author),
                warnings=self.warnings,
                errors=self.errors,
                stats={},
            )

        suffix = path.suffix.lower()
        if suffix in (".md", ".mdx", ".markdown"):
            return self._ingest_markdown(path, title, author)
        elif suffix == ".txt":
            return self._ingest_text(path, title, author)
        elif suffix == ".json":
            return self._ingest_json(path, title, author)
        elif suffix == ".html" or suffix == ".htm":
            return self._ingest_html(path, title, author)
        elif suffix == ".docx":
            return self._ingest_docx(path, title, author)
        else:
            self.errors.append(f"Unsupported file format: {suffix}")
            return IngestionResult(
                manuscript=Manuscript(id="", title=title or path.stem, author=author),
                warnings=self.warnings,
                errors=self.errors,
                stats={},
            )

    def ingest_multiple(self, file_paths: list[str | Path], title: str = "", author: str = "") -> IngestionResult:
        all_blocks: list[ContentBlock] = []
        all_warnings: list[str] = []
        all_errors: list[str] = []
        total_words = 0

        for i, fp in enumerate(file_paths):
            result = self.ingest_file(fp, title if i == 0 else "", author if i == 0 else "")
            all_blocks.extend(result.manuscript.content_blocks)
            all_warnings.extend(result.warnings)
            all_errors.extend(result.errors)
            total_words += result.stats.get("word_count", 0)

        # Re-assign stable order
        for idx, block in enumerate(all_blocks):
            block.order = idx

        manuscript = Manuscript(
            id=f"ms_{uuid.uuid4().hex[:12]}",
            title=title or "Combined Manuscript",
            author=author,
            source_format="mixed",
            source_files=[str(p) for p in file_paths],
            content_blocks=all_blocks,
            metadata={"ingestion_warnings": all_warnings, "ingestion_errors": all_errors},
        )

        return IngestionResult(
            manuscript=manuscript,
            warnings=all_warnings,
            errors=all_errors,
            stats={"word_count": total_words, "block_count": len(all_blocks), "file_count": len(file_paths)},
        )

    def ingest_from_api(self, content: str, format: str, title: str = "", author: str = "") -> IngestionResult:
        if format == "markdown":
            return self._ingest_markdown_content(content, title, author)
        elif format == "text":
            return self._ingest_text_content(content, title, author)
        elif format == "json":
            return self._ingest_json_content(content, title, author)
        elif format == "html":
            return self._ingest_html_content(content, title, author)
        else:
            self.errors.append(f"Unsupported format: {format}")
            return IngestionResult(
                manuscript=Manuscript(id="", title=title, author=author),
                warnings=self.warnings,
                errors=self.errors,
                stats={},
            )

    def _reset_counters(self):
        self.block_counter = 0
        self.warnings = []
        self.errors = []

    def _next_block_id(self) -> str:
        self.block_counter += 1
        return f"cb_{uuid.uuid4().hex[:10]}_{self.block_counter:04d}"

    def _ingest_markdown(self, path: Path, title: str, author: str) -> IngestionResult:
        content = path.read_text(encoding="utf-8")
        return self._ingest_markdown_content(content, title or path.stem, author)

    def _ingest_markdown_content(self, content: str, title: str, author: str) -> IngestionResult:
        self._reset_counters()
        blocks: list[ContentBlock] = []

        lines = content.split("\n")
        in_code_block = False
        code_block_content = []
        code_block_lang = ""
        in_list = False
        list_items: list[str] = []
        current_chapter = 0
        current_section = 0
        current_subsection = 0
        paragraph_buffer: list[str] = []

        def flush_paragraph():
            nonlocal paragraph_buffer
            if paragraph_buffer:
                text = "\n".join(paragraph_buffer).strip()
                if text:
                    blocks.append(self._make_block(
                        text, ContentType.PARAGRAPH, SemanticRole.BODY,
                        current_chapter, current_section, current_subsection
                    ))
                paragraph_buffer = []

        def flush_list():
            nonlocal list_items, in_list
            if list_items:
                for item in list_items:
                    blocks.append(self._make_block(
                        item, ContentType.LIST_ITEM, SemanticRole.BODY,
                        current_chapter, current_section, current_subsection
                    ))
                list_items = []
                in_list = False

        for line in lines:
            stripped = line.strip()

            # Code block handling
            if stripped.startswith("```"):
                flush_paragraph()
                flush_list()
                if not in_code_block:
                    in_code_block = True
                    code_block_lang = stripped[3:].strip()
                    code_block_content = []
                else:
                    in_code_block = False
                    code_text = "\n".join(code_block_content)
                    blocks.append(self._make_block(
                        code_text, ContentType.CODE, SemanticRole.CODE,
                        current_chapter, current_section, current_subsection,
                        metadata={"language": code_block_lang}
                    ))
                    code_block_content = []
                    code_block_lang = ""
                continue

            if in_code_block:
                code_block_content.append(line)
                continue

            # Headings
            heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
            if heading_match:
                flush_paragraph()
                flush_list()
                level = len(heading_match.group(1))
                heading_text = heading_match.group(2).strip()
                if level == 1:
                    current_chapter += 1
                    current_section = 0
                    current_subsection = 0
                    role = SemanticRole.CHAPTER
                elif level == 2:
                    current_section += 1
                    current_subsection = 0
                    role = SemanticRole.SECTION
                elif level == 3:
                    current_subsection += 1
                    role = SemanticRole.SUBSECTION
                else:
                    role = SemanticRole.BODY
                blocks.append(self._make_block(
                    heading_text, ContentType.HEADING, role,
                    current_chapter, current_section, current_subsection,
                    level=level
                ))
                continue

            # List items
            list_match = re.match(r"^[\-\*\+]\s+(.+)$", stripped) or re.match(r"^\d+\.\s+(.+)$", stripped)
            if list_match:
                flush_paragraph()
                in_list = True
                list_items.append(list_match.group(1).strip())
                continue

            # Horizontal rule
            if re.match(r"^[-*_]{3,}$", stripped):
                flush_paragraph()
                flush_list()
                blocks.append(self._make_block(
                    "", ContentType.SECTION_BREAK, SemanticRole.BODY,
                    current_chapter, current_section, current_subsection
                ))
                continue

            # Blockquote
            if stripped.startswith(">"):
                flush_paragraph()
                flush_list()
                quote_text = stripped[1:].strip()
                blocks.append(self._make_block(
                    quote_text, ContentType.QUOTATION, SemanticRole.PULL_QUOTE,
                    current_chapter, current_section, current_subsection
                ))
                continue

            # Table (simplified detection)
            if "|" in stripped and stripped.count("|") >= 2:
                flush_paragraph()
                flush_list()
                # Check if it's a table separator
                if not re.match(r"^\s*\|?\s*[:-]+\s*\|", stripped):
                    blocks.append(self._make_block(
                        stripped, ContentType.TABLE, SemanticRole.TABLE,
                        current_chapter, current_section, current_subsection
                    ))
                continue

            # Image instruction
            img_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
            if img_match:
                flush_paragraph()
                flush_list()
                alt_text = img_match.group(1)
                src = img_match.group(2)
                blocks.append(self._make_block(
                    f"Image: {alt_text} ({src})", ContentType.IMAGE_INSTRUCTION, SemanticRole.FIGURE,
                    current_chapter, current_section, current_subsection,
                    metadata={"alt_text": alt_text, "src": src}
                ))
                continue

            # Regular paragraph content
            if stripped:
                paragraph_buffer.append(stripped)
            else:
                flush_paragraph()

        flush_paragraph()
        flush_list()

        manuscript = Manuscript(
            id=f"ms_{uuid.uuid4().hex[:12]}",
            title=title,
            author=author,
            source_format="markdown",
            source_files=[str(path)] if 'path' in locals() else [],
            content_blocks=blocks,
        )

        return IngestionResult(
            manuscript=manuscript,
            warnings=self.warnings,
            errors=self.errors,
            stats={"word_count": manuscript.total_word_count(), "block_count": len(blocks)},
        )

    def _ingest_text(self, path: Path, title: str, author: str) -> IngestionResult:
        content = path.read_text(encoding="utf-8")
        return self._ingest_text_content(content, title or path.stem, author)

    def _ingest_text_content(self, content: str, title: str, author: str) -> IngestionResult:
        self._reset_counters()
        blocks: list[ContentBlock] = []

        paragraphs = re.split(r"\n\s*\n", content)
        chapter = 0

        for para in paragraphs:
            para = para.strip()
            if not para:
                continue

            # Check if it looks like a heading (short line, no period at end)
            lines = para.split("\n")
            if len(lines) == 1 and len(para) < 100 and not para.endswith("."):
                chapter += 1
                blocks.append(self._make_block(
                    para, ContentType.HEADING, SemanticRole.CHAPTER,
                    chapter, 0, 0, level=1
                ))
            else:
                blocks.append(self._make_block(
                    para, ContentType.PARAGRAPH, SemanticRole.BODY,
                    chapter, 0, 0
                ))

        manuscript = Manuscript(
            id=f"ms_{uuid.uuid4().hex[:12]}",
            title=title,
            author=author,
            source_format="text",
            source_files=[str(path)] if 'path' in locals() else [],
            content_blocks=blocks,
        )

        return IngestionResult(
            manuscript=manuscript,
            warnings=self.warnings,
            errors=self.errors,
            stats={"word_count": manuscript.total_word_count(), "block_count": len(blocks)},
        )

    def _ingest_json(self, path: Path, title: str, author: str) -> IngestionResult:
        content = path.read_text(encoding="utf-8")
        return self._ingest_json_content(content, title or path.stem, author)

    def _ingest_json_content(self, content: str, title: str, author: str) -> IngestionResult:
        self._reset_counters()
        try:
            data = json.loads(content)
        except json.JSONDecodeError as e:
            self.errors.append(f"Invalid JSON: {e}")
            return IngestionResult(
                manuscript=Manuscript(id="", title=title, author=author),
                warnings=self.warnings,
                errors=self.errors,
                stats={},
            )

        blocks: list[ContentBlock] = []

        # Handle different JSON structures
        if isinstance(data, list):
            items = data
        elif isinstance(data, dict) and "blocks" in data:
            items = data["blocks"]
        elif isinstance(data, dict) and "content" in data:
            items = data["content"]
        else:
            items = [data]

        chapter = 0
        section = 0
        for item in items:
            if isinstance(item, str):
                blocks.append(self._make_block(item, ContentType.PARAGRAPH, SemanticRole.BODY, chapter, section, 0))
            elif isinstance(item, dict):
                ctype = item.get("type", "paragraph")
                role = item.get("role", "body")
                try:
                    ct = ContentType(ctype)
                except ValueError:
                    ct = ContentType.PARAGRAPH
                try:
                    sr = SemanticRole(role)
                except ValueError:
                    sr = SemanticRole.BODY

                if ct == ContentType.HEADING and sr == SemanticRole.CHAPTER:
                    chapter += 1
                    section = 0

                blocks.append(self._make_block(
                    item.get("content", ""), ct, sr,
                    chapter, section, 0,
                    level=item.get("level", 0),
                    metadata=item.get("metadata", {})
                ))

        manuscript = Manuscript(
            id=f"ms_{uuid.uuid4().hex[:12]}",
            title=title,
            author=author,
            source_format="json",
            content_blocks=blocks,
        )

        return IngestionResult(
            manuscript=manuscript,
            warnings=self.warnings,
            errors=self.errors,
            stats={"word_count": manuscript.total_word_count(), "block_count": len(blocks)},
        )

    def _ingest_html(self, path: Path, title: str, author: str) -> IngestionResult:
        content = path.read_text(encoding="utf-8")
        return self._ingest_html_content(content, title or path.stem, author)

    def _ingest_html_content(self, content: str, title: str, author: str) -> IngestionResult:
        self._reset_counters()
        blocks: list[ContentBlock] = []

        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(content, "html.parser")
        except ImportError:
            self.errors.append("BeautifulSoup4 required for HTML ingestion")
            return IngestionResult(
                manuscript=Manuscript(id="", title=title, author=author),
                warnings=self.warnings,
                errors=self.errors,
                stats={},
            )

        # Remove script and style elements
        for script in soup(["script", "style"]):
            script.decompose()

        chapter = 0
        section = 0

        for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "blockquote", "pre", "code", "ul", "ol", "table", "img"]):
            if element.name in ("h1", "h2", "h3", "h4", "h5", "h6"):
                level = int(element.name[1])
                text = element.get_text(strip=True)
                if level == 1:
                    chapter += 1
                    section = 0
                    role = SemanticRole.CHAPTER
                elif level == 2:
                    section += 1
                    role = SemanticRole.SECTION
                else:
                    role = SemanticRole.SUBSECTION
                blocks.append(self._make_block(text, ContentType.HEADING, role, chapter, section, 0, level=level))
            elif element.name == "p":
                text = element.get_text(strip=True)
                if text:
                    blocks.append(self._make_block(text, ContentType.PARAGRAPH, SemanticRole.BODY, chapter, section, 0))
            elif element.name == "blockquote":
                text = element.get_text(strip=True)
                if text:
                    blocks.append(self._make_block(text, ContentType.QUOTATION, SemanticRole.PULL_QUOTE, chapter, section, 0))
            elif element.name in ("pre", "code"):
                text = element.get_text()
                if text.strip():
                    blocks.append(self._make_block(text, ContentType.CODE, SemanticRole.CODE, chapter, section, 0))
            elif element.name in ("ul", "ol"):
                for li in element.find_all("li", recursive=False):
                    text = li.get_text(strip=True)
                    if text:
                        blocks.append(self._make_block(text, ContentType.LIST_ITEM, SemanticRole.BODY, chapter, section, 0))
            elif element.name == "table":
                text = element.get_text(strip=True)
                if text:
                    blocks.append(self._make_block(text, ContentType.TABLE, SemanticRole.TABLE, chapter, section, 0))
            elif element.name == "img":
                alt = element.get("alt", "")
                src = element.get("src", "")
                blocks.append(self._make_block(
                    f"Image: {alt} ({src})", ContentType.IMAGE_INSTRUCTION, SemanticRole.FIGURE,
                    chapter, section, 0, metadata={"alt_text": alt, "src": src}
                ))

        manuscript = Manuscript(
            id=f"ms_{uuid.uuid4().hex[:12]}",
            title=title,
            author=author,
            source_format="html",
            content_blocks=blocks,
        )

        return IngestionResult(
            manuscript=manuscript,
            warnings=self.warnings,
            errors=self.errors,
            stats={"word_count": manuscript.total_word_count(), "block_count": len(blocks)},
        )

    def _ingest_docx(self, path: Path, title: str, author: str) -> IngestionResult:
        try:
            from docx import Document
        except ImportError:
            self.errors.append("python-docx required for DOCX ingestion")
            return IngestionResult(
                manuscript=Manuscript(id="", title=title, author=author),
                warnings=self.warnings,
                errors=self.errors,
                stats={},
            )

        self._reset_counters()
        blocks: list[ContentBlock] = []

        doc = Document(path)
        chapter = 0
        section = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style = para.style.name.lower() if para.style else ""

            if "heading" in style or style.startswith("heading"):
                level_match = re.search(r"(\d+)", style)
                level = int(level_match.group(1)) if level_match else 1
                if level == 1:
                    chapter += 1
                    section = 0
                    role = SemanticRole.CHAPTER
                elif level == 2:
                    section += 1
                    role = SemanticRole.SECTION
                else:
                    role = SemanticRole.SUBSECTION
                blocks.append(self._make_block(text, ContentType.HEADING, role, chapter, section, 0, level=level))
            elif "list" in style or para.style.name.startswith("List"):
                blocks.append(self._make_block(text, ContentType.LIST_ITEM, SemanticRole.BODY, chapter, section, 0))
            elif "code" in style:
                blocks.append(self._make_block(text, ContentType.CODE, SemanticRole.CODE, chapter, section, 0))
            elif "quote" in style or "block" in style:
                blocks.append(self._make_block(text, ContentType.QUOTATION, SemanticRole.PULL_QUOTE, chapter, section, 0))
            else:
                blocks.append(self._make_block(text, ContentType.PARAGRAPH, SemanticRole.BODY, chapter, section, 0))

        # Handle tables
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                table_text.append(" | ".join(row_text))
            if table_text:
                blocks.append(self._make_block("\n".join(table_text), ContentType.TABLE, SemanticRole.TABLE, chapter, section, 0))

        manuscript = Manuscript(
            id=f"ms_{uuid.uuid4().hex[:12]}",
            title=title or path.stem,
            author=author,
            source_format="docx",
            source_files=[str(path)],
            content_blocks=blocks,
        )

        return IngestionResult(
            manuscript=manuscript,
            warnings=self.warnings,
            errors=self.errors,
            stats={"word_count": manuscript.total_word_count(), "block_count": len(blocks)},
        )

    def _make_block(
        self,
        content: str,
        content_type: ContentType,
        semantic_role: SemanticRole,
        chapter: int,
        section: int,
        subsection: int,
        level: int = 0,
        metadata: dict[str, Any] | None = None,
    ) -> ContentBlock:
        return ContentBlock(
            id=self._next_block_id(),
            content=content,
            content_type=content_type,
            semantic_role=semantic_role,
            chapter=chapter,
            section=section,
            subsection=subsection,
            order=self.block_counter,
            level=level,
            source_ref=f"ch{chapter}_s{section}_ss{subsection}_{self.block_counter}",
            metadata=metadata or {},
        )